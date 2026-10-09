# -*- coding: utf-8 -*-
"""Inspect and replace the TTF/OTF fonts embedded in Unity files.

    python tools/font.py list     <file>
    python tools/font.py extract  <file> <font name> <out.ttf>
    python tools/font.py replace  <file> <font name> <new.ttf> --out <file> [--apply]
    python tools/font.py coverage <font.ttf> [--charset vietnamese] [--text STR] [--sheet X.xlsx]

TextMeshPro fonts in recent games are usually *dynamic*: the TMP_FontAsset has
an empty glyph table and rasterises glyphs at run time from its source font,
which is a plain `Font` object whose `m_FontData` is the whole TTF. Adding a
script the game lacks (Vietnamese diacritics on a Japanese face) is then a
matter of putting a TTF that has those glyphs into `m_FontData` - no atlas to
rebuild. Check `coverage` first: whatever the TTF lacks falls back to the
TMP fallback chain or draws as a box.

The same font is often duplicated in several files (each UI bundle carries its
own copy), and every copy has to be replaced; `list` each candidate file.

A *static* TMP font (glyphs baked into an atlas texture) is not handled here:
its atlas and glyph table have to be regenerated in the Unity editor.

`coverage --sheet` reads column C (Translation) of a jsonsheet workbook and
reports every character the font cannot draw.
"""
import argparse
import io
import os
import sys

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from unityio import (add_common, backup, load, objects, raw_snapshot, save,   # noqa: E402
                     utf8_stdout, verify_untouched)

VIETNAMESE = ("aàáảãạăằắẳẵặâầấẩẫậeèéẻẽẹêềếểễệiìíỉĩịoòóỏõọôồốổỗộơờớởỡợ"
              "uùúủũụưừứửữựyỳýỷỹỵđ")
LATIN = "".join(chr(c) for c in range(0x21, 0x7F))
CHARSETS = {"vietnamese": VIETNAMESE + VIETNAMESE.upper(), "latin": LATIN}


def fonts(env, name=None):
    rows = [(o, d) for o, d in objects(env, "Font") if name is None or d.m_Name == name]
    return rows


def font_bytes(d):
    return bytes(d.m_FontData or b"")


def describe(data):
    if not data:
        return "no embedded data (OS / dynamic reference)"
    try:
        from fontTools.ttLib import TTFont
        tt = TTFont(io.BytesIO(data), lazy=True)
        name = tt["name"].getDebugName(4) or "?"
        n = len(tt.getBestCmap() or {})
        return "%s, %d mapped characters" % (name, n)
    except Exception as e:
        return "unreadable font data (%s)" % e


def cmd_list(a):
    env = load(a.file, a.unity_version)
    rows = fonts(env)
    for o, d in rows:
        data = font_bytes(d)
        print("  %20d  %-32s %10d bytes  %s" % (o.path_id, d.m_Name, len(data), describe(data)))
    print("%d Font objects" % len(rows))


def one(env, name):
    rows = fonts(env, name)
    if len(rows) != 1:
        raise SystemExit("%d Font objects named %r - `list` shows what is there" % (len(rows), name))
    return rows[0]


def cmd_extract(a):
    _o, d = one(load(a.file, a.unity_version), a.name)
    data = font_bytes(d)
    if not data:
        raise SystemExit("%s has no embedded font data" % a.name)
    with open(a.out, "wb") as f:
        f.write(data)
    print("wrote %s (%d bytes): %s" % (a.out, len(data), describe(data)))


def cmd_replace(a):
    env = load(a.file, a.unity_version)
    before = raw_snapshot(env)
    o, d = one(env, a.name)
    data = open(a.ttf, "rb").read()
    info = describe(data)
    if info.startswith("unreadable"):
        raise SystemExit("%s: %s" % (a.ttf, info))
    print("%s: %d -> %d bytes" % (a.name, len(font_bytes(d)), len(data)))
    print("   old: %s" % describe(font_bytes(d)))
    print("   new: %s" % info)
    d.m_FontData = list(data) if isinstance(d.m_FontData, list) else data
    d.save()
    if not a.apply:
        print("dry run - add --apply to write %s" % a.out)
        return
    if os.path.abspath(a.out) == os.path.abspath(a.file):
        print("backup -> %s" % backup(a.file))
    size = save(env, a.out, a.packer)
    chk = verify_untouched(before, a.out, {o.path_id}, a.unity_version)
    got = [font_bytes(x.read()) for x in chk.objects if x.path_id == o.path_id][0]
    if got != data:
        raise SystemExit("the font read back differs from %s" % a.ttf)
    print("wrote %s (%d bytes); font reads back identical, other objects untouched" % (a.out, size))


def cmd_coverage(a):
    from fontTools.ttLib import TTFont
    cmap = TTFont(a.font, lazy=True).getBestCmap() or {}
    text = a.text or ""
    for cs in a.charset or []:
        text += CHARSETS[cs]
    if a.sheet:
        from openpyxl import load_workbook
        wb = load_workbook(a.sheet, read_only=True)
        for ws in wb.worksheets:
            for r in ws.iter_rows(min_row=2, min_col=3, max_col=3, values_only=True):
                if r[0]:
                    text += str(r[0])
    chars = sorted({c for c in text if not c.isspace()})
    if not chars:
        raise SystemExit("nothing to check - pass --charset, --text or --sheet")
    missing = [c for c in chars if ord(c) not in cmap]
    print("%s: %d mapped characters; checked %d distinct, %d missing"
          % (os.path.basename(a.font), len(cmap), len(chars), len(missing)))
    if missing:
        print("  missing: " + " ".join("%s(U+%04X)" % (c, ord(c)) for c in missing))
    sys.exit(1 if missing else 0)


def main():
    utf8_stdout()
    ap = argparse.ArgumentParser(description=__doc__.split("\n")[0])
    sub = ap.add_subparsers(dest="cmd", required=True)
    p = sub.add_parser("list")
    p.add_argument("file")
    add_common(p)
    p = sub.add_parser("extract")
    p.add_argument("file")
    p.add_argument("name")
    p.add_argument("out")
    add_common(p)
    p = sub.add_parser("replace")
    p.add_argument("file")
    p.add_argument("name")
    p.add_argument("ttf")
    p.add_argument("--out", required=True, help="output file (may be the input itself)")
    p.add_argument("--packer", default="auto", choices=["auto", "lz4", "lzma", "none", "original"])
    p.add_argument("--apply", action="store_true", help="write (default: dry run)")
    add_common(p)
    p = sub.add_parser("coverage")
    p.add_argument("font")
    p.add_argument("--charset", action="append", choices=sorted(CHARSETS))
    p.add_argument("--text", help="characters to check")
    p.add_argument("--sheet", help="workbook whose column C to check")
    a = ap.parse_args()
    {"list": cmd_list, "extract": cmd_extract, "replace": cmd_replace, "coverage": cmd_coverage}[a.cmd](a)


if __name__ == "__main__":
    main()
