# -*- coding: utf-8 -*-
"""Translate the strings of JSON TextAssets through a spreadsheet.

    python tools/jsonsheet.py export <file> <sheet.xlsx> [--name GLOB] [--key REGEX] [--path REGEX]
    python tools/jsonsheet.py apply  <file> <sheet.xlsx> --out <file> [--token REGEX] [--apply]

Unity games commonly keep dialogue and UI tables as JSON TextAssets (story text
in one large scenario asset, menus and glossaries in a `json` bundle). `export`
writes every string value to a sheet with columns ID | Source | Translation,
one row per value; `apply` writes the Translation column back.

The ID is `<asset name>#<JSON pointer>`, e.g. `DictionaryData#/data/16/text/jp`
(RFC 6901: `~1` stands for `/` and `~0` for `~` inside a key), so a row keeps
its address however the sheet is sorted or filtered.

Filters for `export`:
  --key REGEX   only values whose own key matches, e.g. `^(jp|text)$` for a game
                that stores languages side by side and only reads one slot
  --path REGEX  only pointers matching, e.g. `^/data/\\d+/text/`
Values that are empty, or contain no letter at all, are skipped unless --all.

`apply` edits the raw JSON text at each value's exact position rather than
re-serialising the whole document, so key order, spacing, escapes and the
BOM are kept and only translated values change. Before writing it checks:
  - the row's Source still equals the value in the file (a stale sheet would
    otherwise overwrite a value that has moved on), unless --force;
  - every --token REGEX (e.g. `\\[[^\\]]*\\]` for [name] tags) occurs the same
    number of times in Source and Translation;
  - the result parses and differs from the original exactly at those values.
Line breaks are real newlines in the sheet and `\\n` in the JSON. An empty
Translation cell means "not translated"; write `<empty>` to blank a value.
"""
import argparse
import json
import os
import re
import sys

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from unityio import (add_common, backup, load, objects, raw_snapshot, save,   # noqa: E402
                     text_of, utf8_stdout, verify_untouched)

BOM = "﻿"
STRING_RX = re.compile(r'"(?:[^"\\]|\\.)*"', re.S)
SCALAR_RX = re.compile(r'-?\d+(?:\.\d+)?(?:[eE][+-]?\d+)?|true|false|null')
WS_RX = re.compile(r'\s*')
LETTER_RX = re.compile(r'[^\W\d_]')
CELL_LIMIT = 32767
EMPTY = "<empty>"


def esc(key):
    return key.replace("~", "~0").replace("/", "~1")


def scan(text):
    """{pointer: (start, end)} of every string value, positions in `text`."""
    spans = {}
    pos = [0]

    def ws():
        pos[0] = WS_RX.match(text, pos[0]).end()

    def string():
        m = STRING_RX.match(text, pos[0])
        if not m:
            raise ValueError("expected a string at %d" % pos[0])
        pos[0] = m.end()
        return m.start(), m.end()

    def value(ptr):
        ws()
        c = text[pos[0]:pos[0] + 1]
        if c == "{":
            pos[0] += 1
            ws()
            if text.startswith("}", pos[0]):
                pos[0] += 1
                return
            while True:
                ws()
                s, e = string()
                key = json.loads(text[s:e])
                ws()
                if not text.startswith(":", pos[0]):
                    raise ValueError("expected ':' at %d" % pos[0])
                pos[0] += 1
                value("%s/%s" % (ptr, esc(key)))
                ws()
                if text.startswith(",", pos[0]):
                    pos[0] += 1
                elif text.startswith("}", pos[0]):
                    pos[0] += 1
                    return
                else:
                    raise ValueError("expected ',' or '}' at %d" % pos[0])
        elif c == "[":
            pos[0] += 1
            ws()
            if text.startswith("]", pos[0]):
                pos[0] += 1
                return
            i = 0
            while True:
                value("%s/%d" % (ptr, i))
                i += 1
                ws()
                if text.startswith(",", pos[0]):
                    pos[0] += 1
                elif text.startswith("]", pos[0]):
                    pos[0] += 1
                    return
                else:
                    raise ValueError("expected ',' or ']' at %d" % pos[0])
        elif c == '"':
            if ptr in spans:
                raise ValueError("duplicate key at %s" % ptr)
            spans[ptr] = string()
        else:
            m = SCALAR_RX.match(text, pos[0])
            if not m:
                raise ValueError("unexpected %r at %d" % (c, pos[0]))
            pos[0] = m.end()

    start = len(BOM) if text.startswith(BOM) else 0
    pos[0] = start
    value("")
    ws()
    if pos[0] != len(text):
        raise ValueError("trailing data at %d" % pos[0])
    return spans


def json_assets(env, glob):
    out = []
    for o, d in objects(env, "TextAsset", glob):
        text = text_of(d)
        try:
            spans = scan(text)
        except ValueError:
            continue
        out.append((o, d, text, spans))
    return out


def cmd_export(a):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font
    env = load(a.file, a.unity_version)
    key_rx = re.compile(a.key) if a.key else None
    path_rx = re.compile(a.path) if a.path else None
    wb = Workbook()
    ws = wb.active
    ws.title = "text"
    ws.append(["ID", "Source", "Translation"])
    for c in ws[1]:
        c.font = Font(bold=True)
    ws.freeze_panes = "A2"
    rows = too_long = 0
    for o, d, text, spans in json_assets(env, a.name):
        n = 0
        for ptr, (s, e) in spans.items():
            last = ptr.rsplit("/", 1)[-1].replace("~1", "/").replace("~0", "~")
            if key_rx and not key_rx.search(last):
                continue
            if path_rx and not path_rx.search(ptr):
                continue
            val = json.loads(text[s:e])
            if not a.all and not LETTER_RX.search(val):
                continue
            if len(val) > CELL_LIMIT:
                too_long += 1
                continue
            ws.append(["%s#%s" % (d.m_Name, ptr), val, None])
            n += 1
        rows += n
        print("  %-36s %6d strings" % (d.m_Name, n))
    for col, width in (("A", 42), ("B", 70), ("C", 70)):
        ws.column_dimensions[col].width = width
    for row in ws.iter_rows(min_row=2, min_col=2, max_col=3):
        for c in row:
            c.alignment = Alignment(wrap_text=True, vertical="top")
    wb.save(a.sheet)
    print("%d rows -> %s" % (rows, a.sheet))
    if too_long:
        print("WARNING: %d values exceed Excel's %d-character cell limit and were left out"
              % (too_long, CELL_LIMIT))


def read_sheet(path, sheet):
    from openpyxl import load_workbook
    wb = load_workbook(path, read_only=True)
    ws = wb[sheet] if sheet else wb.worksheets[0]
    rows = {}
    for i, r in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
        if not r or not r[0]:
            continue
        rid = str(r[0]).strip()
        src = "" if len(r) < 2 or r[1] is None else str(r[1])
        tr = None if len(r) < 3 or r[2] is None or str(r[2]) == "" else str(r[2])
        if tr is None:
            continue
        if rid in rows:
            raise SystemExit("row %d repeats ID %s" % (i, rid))
        tr = "" if tr == EMPTY else tr.replace("\r\n", "\n")
        rows[rid] = (i, src.replace("\r\n", "\n"), tr)
    return rows


def encode_like(original_literal, value):
    """JSON literal for `value`, escaping non-ASCII only if the file does."""
    ascii_only = "\\u" in original_literal and not any(ord(ch) > 127 for ch in original_literal)
    return json.dumps(value, ensure_ascii=ascii_only)


def cmd_apply(a):
    env = load(a.file, a.unity_version)
    before = raw_snapshot(env)
    rows = read_sheet(a.sheet, a.sheet_name)
    tokens = [re.compile(t) for t in a.token or []]
    by_asset = {}
    for rid, entry in rows.items():
        if "#" not in rid:
            raise SystemExit("row %d: ID %r is not <asset>#<pointer>" % (entry[0], rid))
        asset, ptr = rid.split("#", 1)
        by_asset.setdefault(asset, {})[ptr] = entry
    assets = {d.m_Name: (o, d, text, spans) for o, d, text, spans in json_assets(env, None)}
    problems, changed, written = [], set(), 0
    for asset, entries in sorted(by_asset.items()):
        if asset not in assets:
            problems.append("no JSON TextAsset named %s (%d rows)" % (asset, len(entries)))
            continue
        o, d, text, spans = assets[asset]
        edits = []
        for ptr, (row, src, tr) in entries.items():
            if ptr not in spans:
                problems.append("row %d: %s#%s is not a string value in the file" % (row, asset, ptr))
                continue
            s, e = spans[ptr]
            cur = json.loads(text[s:e])
            if cur != src and not a.force:
                problems.append("row %d: Source no longer matches the file at %s#%s" % (row, asset, ptr))
                continue
            bad = [t.pattern for t in tokens if len(t.findall(src)) != len(t.findall(tr))]
            if bad:
                problems.append("row %d: token count differs for %s" % (row, ", ".join(bad)))
                continue
            if tr != cur:
                edits.append((s, e, ptr, tr))
        if not edits:
            continue
        out = text
        for s, e, ptr, tr in sorted(edits, reverse=True):
            out = out[:s] + encode_like(text[s:e], tr) + out[e:]
        new_spans = scan(out)
        if new_spans.keys() != spans.keys():
            raise SystemExit("%s: structure changed while applying - aborting" % asset)
        want = {ptr: tr for _s, _e, ptr, tr in edits}
        for ptr, (s, e) in new_spans.items():
            got = json.loads(out[s:e])
            old = json.loads(text[spans[ptr][0]:spans[ptr][1]])
            if got != want.get(ptr, old):
                raise SystemExit("%s: value at %s is not what was intended" % (asset, ptr))
        d.m_Script = out
        d.save()
        changed.add(o.path_id)
        written += len(edits)
        print("  %-36s %6d values translated" % (asset, len(edits)))
    for p in problems[:50]:
        print("  SKIP " + p)
    if len(problems) > 50:
        print("  ... %d more skipped" % (len(problems) - 50))
    print("%d values to write, %d rows skipped" % (written, len(problems)))
    if not changed:
        return
    if not a.apply:
        print("dry run - add --apply to write %s" % a.out)
        return
    if os.path.abspath(a.out) == os.path.abspath(a.file):
        print("backup -> %s" % backup(a.file))
    size = save(env, a.out, a.packer)
    verify_untouched(before, a.out, changed, a.unity_version)
    print("wrote %s (%d bytes); every other object byte-identical" % (a.out, size))


def main():
    utf8_stdout()
    ap = argparse.ArgumentParser(description=__doc__.split("\n")[0])
    sub = ap.add_subparsers(dest="cmd", required=True)
    p = sub.add_parser("export")
    p.add_argument("file")
    p.add_argument("sheet")
    p.add_argument("--name", help="only TextAssets whose name matches this glob")
    p.add_argument("--key", help="only values whose key matches this regex")
    p.add_argument("--path", help="only values whose pointer matches this regex")
    p.add_argument("--all", action="store_true", help="keep empty and letter-less values")
    add_common(p)
    p = sub.add_parser("apply")
    p.add_argument("file")
    p.add_argument("sheet")
    p.add_argument("--out", required=True, help="output file (may be the input itself)")
    p.add_argument("--sheet-name", help="worksheet to read (default: the first)")
    p.add_argument("--token", action="append", help="regex that must occur equally often in "
                                                    "Source and Translation; repeatable")
    p.add_argument("--force", action="store_true", help="write even where Source no longer matches")
    p.add_argument("--packer", default="auto", choices=["auto", "lz4", "lzma", "none", "original"])
    p.add_argument("--apply", action="store_true", help="write (default: dry run)")
    add_common(p)
    a = ap.parse_args()
    (cmd_export if a.cmd == "export" else cmd_apply)(a)


if __name__ == "__main__":
    main()
