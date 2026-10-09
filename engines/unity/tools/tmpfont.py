# -*- coding: utf-8 -*-
"""Add characters to a *static* TextMeshPro font asset, without the Unity Editor.

    python tools/tmpfont.py info   <file> [--charset vietnamese] [--text STR] [--sheet X.xlsx]
    python tools/tmpfont.py verify <file> <font asset> <source.ttf>
    python tools/tmpfont.py add    <file> <font asset> <font.ttf> [--charset vietnamese]
                                   [--text STR] [--sheet X.xlsx] [--replace] --out <file>
                                   [--preview P.png] [--apply]

A static TMP font asset (`m_AtlasPopulationMode` 0) draws only the glyphs baked
into its SDF atlas texture; anything else falls back or shows as a box, and
swapping the source TTF changes nothing. `add` bakes new glyphs into the free
space of the atlas exactly the way TMP's own generator does:

  - metrics at the asset's sampling point size from the outline's control box,
    rounded to 1/64 px (FreeType's 26.6) - reproduces all 7,125 stock glyphs of
    UNLOGICAL's font to within 1/64 px;
  - the glyph rect is that box snapped outward to whole pixels; the reserved
    area is the rect grown by padding+1 on the left/bottom and padding on the
    right/top, as in the stock `m_UsedGlyphRects`;
  - the tile is a signed distance field: 0.5 on the outline, falling by
    1/(2 * _GradientScale) per pixel, the scale TMP's shader decodes with (read
    from the asset's material);
  - glyphs are packed into `m_FreeGlyphRects` (MaxRects, best short side fit),
    and the free and used rect lists are kept consistent so the asset can still
    be extended in the Editor later.

`verify` is the positive control: it rebuilds a sample of the asset's existing
glyphs from a TTF and compares metrics and atlas pixels. Run it with the
asset's own source font first - stock UNLOGICAL gives 0 metric errors and a
mean pixel error under 2/255. A large error means the generator does not match
this asset (different TMP version or render mode), and `add` should not be
trusted for it.

The new glyphs may come from a different font than the original (a Japanese
face has no Vietnamese); they are scaled to the same point size and baseline.
`--preview` renders a sample line through the updated atlas so the result can
be looked at before it goes into the game.
"""
import argparse
import math
import os
import sys

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from unityio import (add_common, backup, load, raw_snapshot, save,   # noqa: E402
                     utf8_stdout, verify_untouched)

VIETNAMESE = ("aàáảãạăằắẳẵặâầấẩẫậeèéẻẽẹêềếểễệiìíỉĩịoòóỏõọôồốổỗộơờớởỡợ"
              "uùúủũụưừứửữựyỳýỷỹỵđ")
LATIN = "".join(chr(c) for c in range(0x21, 0x7F))
CHARSETS = {"vietnamese": VIETNAMESE + VIETNAMESE.upper(), "latin": LATIN}
SUPERSAMPLE = 8


# ------------------------------------------------------------------ the asset

def font_assets(env):
    out = []
    for o in env.objects:
        if o.type.name != "MonoBehaviour":
            continue
        try:
            t = o.read_typetree()
        except Exception:
            continue
        if isinstance(t, dict) and "m_CharacterTable" in t and "m_GlyphTable" in t:
            out.append((o, t))
    return out


def find_asset(env, name):
    rows = [(o, t) for o, t in font_assets(env) if t["m_Name"] == name]
    if len(rows) != 1:
        names = ", ".join(t["m_Name"] for _o, t in font_assets(env)) or "none"
        raise SystemExit("%d TMP font assets named %r (present: %s)" % (len(rows), name, names))
    return rows[0]


def gradient_scale(objs, t):
    mat = objs.get(t["m_Material"]["m_PathID"])
    if mat is not None:
        for k, v in mat.read_typetree()["m_SavedProperties"]["m_Floats"]:
            if k == "_GradientScale":
                return float(v)
    return float(t["m_AtlasPadding"] + 1)


def wanted_chars(a):
    text = a.text or ""
    for cs in a.charset or []:
        text += CHARSETS[cs]
    if a.sheet:
        from openpyxl import load_workbook
        for ws in load_workbook(a.sheet, read_only=True).worksheets:
            for r in ws.iter_rows(min_row=2, min_col=3, max_col=3, values_only=True):
                if r[0]:
                    text += str(r[0])
    return sorted({c for c in text if c not in "\r\n\t"})


# ------------------------------------------------------------------ geometry

class Source:
    """A TTF/OTF at the asset's sampling size."""

    def __init__(self, path, point_size, face_index=0):
        from fontTools.ttLib import TTFont
        from PIL import ImageFont
        self.tt = TTFont(path, fontNumber=face_index, lazy=True)
        self.glyphs = self.tt.getGlyphSet()
        self.cmap = self.tt.getBestCmap() or {}
        self.order = self.tt.getGlyphOrder()
        self.scale = point_size / self.tt["head"].unitsPerEm
        self.family = self.tt["name"].getDebugName(1) or ""
        self.pil = ImageFont.truetype(path, point_size * SUPERSAMPLE, index=face_index)

    def has(self, ch):
        return ord(ch) in self.cmap

    def glyph_id(self, ch):
        return self.order.index(self.cmap[ord(ch)])

    def metrics(self, ch):
        """TMP metrics (1/64 px) and the integer rect size, or rect None when blank."""
        from fontTools.pens.boundsPen import ControlBoundsPen
        g = self.glyphs[self.cmap[ord(ch)]]
        q = lambda v: round(v * self.scale * 64) / 64   # noqa: E731
        pen = ControlBoundsPen(self.glyphs)
        g.draw(pen)
        adv = q(g.width)
        if pen.bounds is None:
            return {"m_Width": 0.0, "m_Height": 0.0, "m_HorizontalBearingX": 0.0,
                    "m_HorizontalBearingY": 0.0, "m_HorizontalAdvance": adv}, None
        xmin, ymin, xmax, ymax = pen.bounds
        m = {"m_Width": q(xmax - xmin), "m_Height": q(ymax - ymin), "m_HorizontalBearingX": q(xmin),
             "m_HorizontalBearingY": q(ymax), "m_HorizontalAdvance": adv}
        bx, by = m["m_HorizontalBearingX"], m["m_HorizontalBearingY"]
        w = math.ceil(bx + m["m_Width"]) - math.floor(bx)
        h = math.ceil(by) - math.floor(by - m["m_Height"])
        return m, (w, h)

    def sdf(self, ch, m, size, pad, spread):
        """Tile of (w + 2 pad) x (h + 2 pad), rows top-down, values 0..255."""
        import numpy as np
        from PIL import Image, ImageDraw
        w, h = size[0] + 2 * pad, size[1] + 2 * pad
        ss = SUPERSAMPLE
        img = Image.new("L", (w * ss, h * ss), 0)
        x = (pad - math.floor(m["m_HorizontalBearingX"])) * ss
        y = (pad + math.ceil(m["m_HorizontalBearingY"])) * ss
        ImageDraw.Draw(img).text((x, y), ch, font=self.pil, fill=255, anchor="ls")
        inside = np.array(img) >= 128
        p = np.pad(inside, 1)
        edge = inside & ~(p[:-2, 1:-1] & p[2:, 1:-1] & p[1:-1, :-2] & p[1:-1, 2:])
        ey, ex = np.nonzero(edge)
        if not len(ex):
            return np.zeros((h, w), np.uint8)
        pts = np.stack([(ex + 0.5) / ss, (ey + 0.5) / ss], 1)
        yy, xx = np.mgrid[0:h, 0:w]
        cen = np.stack([xx.ravel() + 0.5, yy.ravel() + 0.5], 1)
        best = np.full(len(cen), np.inf)
        for i in range(0, len(pts), 1024):
            d2 = ((cen[:, None, :] - pts[None, i:i + 1024, :]) ** 2).sum(-1)
            best = np.minimum(best, d2.min(1))
        dist = np.sqrt(best).reshape(h, w)
        cover = inside.reshape(h, ss, w, ss).mean((1, 3))
        signed = np.where(cover >= 0.5, dist, -dist)
        return np.clip(np.rint((0.5 + signed / (2 * spread)) * 255), 0, 255).astype(np.uint8)


# ------------------------------------------------------------------ packing

def rect(r):
    return r["m_X"], r["m_Y"], r["m_Width"], r["m_Height"]


def as_rect(x, y, w, h):
    return {"m_X": int(x), "m_Y": int(y), "m_Width": int(w), "m_Height": int(h)}


def place(free, w, h):
    """Best short side fit; returns (x, y) or None. `free` holds (x, y, w, h)."""
    best = None
    for fx, fy, fw, fh in free:
        if fw >= w and fh >= h:
            key = (min(fw - w, fh - h), max(fw - w, fh - h))
            if best is None or key < best[0]:
                best = (key, fx, fy)
    return None if best is None else (best[1], best[2])


def occupy(free, used):
    """MaxRects split of every free rect `used` overlaps, then prune contained ones."""
    ux, uy, uw, uh = used
    out = []
    for f in free:
        fx, fy, fw, fh = f
        if ux >= fx + fw or ux + uw <= fx or uy >= fy + fh or uy + uh <= fy:
            out.append(f)
            continue
        if ux > fx:
            out.append((fx, fy, ux - fx, fh))
        if ux + uw < fx + fw:
            out.append((ux + uw, fy, fx + fw - ux - uw, fh))
        if uy > fy:
            out.append((fx, fy, fw, uy - fy))
        if uy + uh < fy + fh:
            out.append((fx, uy + uh, fw, fy + fh - uy - uh))
    out = list(dict.fromkeys(r for r in out if r[2] > 0 and r[3] > 0))
    keep = []
    for i, a in enumerate(out):
        ax, ay, aw, ah = a
        if not any(j != i and b[0] <= ax and b[1] <= ay and b[0] + b[2] >= ax + aw
                   and b[1] + b[3] >= ay + ah and (b != a or j < i) for j, b in enumerate(out)):
            keep.append(a)
    return keep


# ------------------------------------------------------------------ commands

def atlas_alpha(objs, t):
    import numpy as np
    texs = t["m_AtlasTextures"]
    if not texs:
        raise SystemExit("the asset has no atlas texture")
    tex_obj = objs[texs[0]["m_PathID"]]
    tex = tex_obj.read()
    img = tex.image
    return tex_obj, tex, img, np.array(img)


def cmd_info(a):
    env = load(a.file, a.unity_version)
    objs = {o.path_id: o for o in env.objects}
    want = wanted_chars(a) if (a.text or a.charset or a.sheet) else []
    rows = font_assets(env)
    for o, t in rows:
        mode = {0: "static", 1: "dynamic", 2: "dynamic OS"}.get(t.get("m_AtlasPopulationMode"), "?")
        free = sum(r["m_Width"] * r["m_Height"] for r in t.get("m_FreeGlyphRects", []))
        print("%20d  %s" % (o.path_id, t["m_Name"]))
        print("      %s, %d glyphs, %d characters, atlas %dx%d, padding %d, gradient scale %g, "
              "%d pt, free rects %d (%.1f Mpx)"
              % (mode, len(t["m_GlyphTable"]), len(t["m_CharacterTable"]), t["m_AtlasWidth"],
                 t["m_AtlasHeight"], t["m_AtlasPadding"], gradient_scale(objs, t),
                 t["m_FaceInfo"]["m_PointSize"], len(t.get("m_FreeGlyphRects", [])), free / 1e6))
        if want:
            have = {c["m_Unicode"] for c in t["m_CharacterTable"]}
            missing = [c for c in want if ord(c) not in have and not c.isspace()]
            print("      missing %d of %d: %s" % (len(missing), len(want), "".join(missing)[:200]))
    print("%d TMP font assets" % len(rows))


def cmd_verify(a):
    import numpy as np
    env = load(a.file, a.unity_version)
    objs = {o.path_id: o for o in env.objects}
    _o, t = find_asset(env, a.asset)
    pt, pad = t["m_FaceInfo"]["m_PointSize"], t["m_AtlasPadding"]
    spread = gradient_scale(objs, t)
    src = Source(a.ttf, pt, a.face_index)
    _to, _tex, _img, arr = atlas_alpha(objs, t)
    alpha = arr[..., 3] if arr.ndim == 3 else arr
    H = alpha.shape[0]
    glyphs = {g["m_Index"]: g for g in t["m_GlyphTable"]}
    chars = [c for c in t["m_CharacterTable"] if src.has(chr(c["m_Unicode"]))]
    step = max(1, len(chars) // a.sample)
    worst_metric, errs, n = 0.0, [], 0
    for c in chars[::step]:
        ch = chr(c["m_Unicode"])
        g = glyphs[c["m_GlyphIndex"]]
        m, size = src.metrics(ch)
        worst_metric = max(worst_metric, max(abs(m[k] - g["m_Metrics"][k]) for k in m))
        x, y, w, h = rect(g["m_GlyphRect"])
        n += 1
        if size is None or w == 0:
            continue
        if size != (w, h):
            worst_metric = max(worst_metric, 1.0)
            continue
        tile = src.sdf(ch, m, size, pad, spread).astype(float)
        top = H - (y - pad + h + 2 * pad)
        stock = alpha[top:top + h + 2 * pad, x - pad:x + w + pad].astype(float)
        errs.append(np.abs(tile - stock).mean())
    if not n:
        raise SystemExit("the TTF has none of the asset's characters - not its source font")
    print("%s vs %s: %d glyphs compared" % (a.asset, os.path.basename(a.ttf), n))
    print("   worst metric error %.4f px (1/64 = 0.0156)" % worst_metric)
    if errs:
        print("   mean pixel error %.2f / 255 over %d same-size tiles, worst glyph %.2f"
              % (np.mean(errs), len(errs), np.max(errs)))
    else:
        print("   no glyph came out the same size as in the atlas")
    ok = worst_metric <= 1 / 64 + 1e-9 and bool(errs) and np.mean(errs) < 4
    print("   %s" % ("MATCHES - `add` reproduces this asset's generator"
                     if ok else "DOES NOT MATCH - wrong source font, or a generator `add` does not reproduce"))
    sys.exit(0 if ok else 1)


def cmd_add(a):
    import numpy as np
    from PIL import Image
    env = load(a.file, a.unity_version)
    before = raw_snapshot(env)
    objs = {o.path_id: o for o in env.objects}
    obj, t = find_asset(env, a.asset)
    if t.get("m_AtlasPopulationMode", 0) != 0:
        raise SystemExit("%s is dynamic - put a TTF with the glyphs into its source Font "
                         "instead (font.py replace)" % a.asset)
    if len(t["m_AtlasTextures"]) != 1:
        raise SystemExit("multi-atlas font assets are not supported")
    pt, pad = t["m_FaceInfo"]["m_PointSize"], t["m_AtlasPadding"]
    spread = gradient_scale(objs, t)
    src = Source(a.ttf, pt, a.face_index)
    have = {c["m_Unicode"] for c in t["m_CharacterTable"]}
    if a.replace:
        # Re-bake characters the asset already has, so a word mixing new and old
        # letters is drawn in one typeface. The old tiles stay in the atlas, unused.
        todo = [c for c in wanted_chars(a) if c != " " or ord(c) not in have]
    else:
        todo = [c for c in wanted_chars(a) if ord(c) not in have]
    lacking = [c for c in todo if not src.has(c)]
    todo = [c for c in todo if src.has(c)]
    if lacking:
        print("not in %s, skipped: %s" % (os.path.basename(a.ttf), "".join(lacking)))
    if not todo:
        print("nothing to add - every requested character is already in %s" % a.asset)
        return
    tex_obj, tex, img, arr = atlas_alpha(objs, t)
    alpha = arr[..., 3].copy() if arr.ndim == 3 else arr.copy()
    H, W = alpha.shape
    if (W, H) != (t["m_AtlasWidth"], t["m_AtlasHeight"]):
        raise SystemExit("atlas texture is %dx%d, asset says %dx%d" % (W, H, t["m_AtlasWidth"], t["m_AtlasHeight"]))
    free = [rect(r) for r in t["m_FreeGlyphRects"]]
    used = [rect(r) for r in t["m_UsedGlyphRects"]]
    used_ids = {g["m_Index"] for g in t["m_GlyphTable"]}
    same_font = src.family and src.family == t["m_FaceInfo"]["m_FamilyName"]
    next_id = max(used_ids | set([len(src.order)])) + 1
    new_glyphs, new_chars, placed = [], [], 0
    print("%s: %g pt, padding %d, gradient scale %g; adding %d characters from %s"
          % (a.asset, pt, pad, spread, len(todo), src.family or a.ttf))
    for ch in sorted(todo, key=lambda c: -(src.metrics(c)[1] or (0, 0))[1]):
        m, size = src.metrics(ch)
        gid = src.glyph_id(ch) if same_font and src.glyph_id(ch) not in used_ids else next_id
        if gid == next_id:
            next_id += 1
        used_ids.add(gid)
        if size is None:
            gr = as_rect(0, 0, 0, 0)
        else:
            need = (size[0] + 2 * pad + 1, size[1] + 2 * pad + 1)
            spot = place(free, *need)
            if spot is None:
                raise SystemExit("the atlas has no free space left for %r (%dx%d); add fewer "
                                 "characters or rebuild the atlas larger in the Editor" % (ch, need[0], need[1]))
            ux, uy = spot
            gr = as_rect(ux + pad + 1, uy + pad + 1, size[0], size[1])
            tile = src.sdf(ch, m, size, pad, spread)
            x0, y0 = gr["m_X"] - pad, gr["m_Y"] - pad
            top = H - (y0 + tile.shape[0])
            region = alpha[top:top + tile.shape[0], x0:x0 + tile.shape[1]]
            if region.shape != tile.shape or region.any():
                raise SystemExit("free rect for %r is not empty in the texture - the free list "
                                 "does not match the atlas; refusing to draw over glyphs" % ch)
            alpha[top:top + tile.shape[0], x0:x0 + tile.shape[1]] = tile
            free = occupy(free, (ux, uy) + need)
            used.append((ux, uy) + need)
            placed += 1
        new_glyphs.append({"m_Index": gid, "m_Metrics": m, "m_GlyphRect": gr, "m_Scale": 1.0,
                           "m_AtlasIndex": 0, "m_ClassDefinitionType": 0})
        new_chars.append({"m_ElementType": 1, "m_Unicode": ord(ch), "m_GlyphIndex": gid, "m_Scale": 1.0})
    t["m_GlyphTable"] = sorted(t["m_GlyphTable"] + new_glyphs, key=lambda g: g["m_Index"])
    replaced = {c["m_Unicode"] for c in new_chars}
    t["m_CharacterTable"] = sorted([c for c in t["m_CharacterTable"] if c["m_Unicode"] not in replaced]
                                   + new_chars, key=lambda c: c["m_Unicode"])
    t["m_UsedGlyphRects"] = [as_rect(*r) for r in used]
    t["m_FreeGlyphRects"] = [as_rect(*r) for r in free]
    print("baked %d glyphs (+%d blank), %d of them replacing existing characters; %d free rects left"
          % (placed, len(new_glyphs) - placed, len(replaced & have), len(free)))
    if a.preview:
        preview(a.preview, a.preview_text or "".join(todo[:40]), t, alpha, pad)
        print("preview -> %s" % a.preview)
    if not a.apply:
        print("dry run - add --apply to write %s" % a.out)
        return
    if arr.ndim == 3:
        arr[..., 3] = alpha
        out_img = Image.fromarray(arr, "RGBA")
    else:
        out_img = Image.fromarray(alpha, "L")
    tex.image = out_img
    tex.save()
    obj.save_typetree(t)
    if os.path.abspath(a.out) == os.path.abspath(a.file):
        print("backup -> %s" % backup(a.file))
    size_b = save(env, a.out, a.packer)
    chk = verify_untouched(before, a.out, {obj.path_id, tex_obj.path_id}, a.unity_version)
    back = {o.path_id: o for o in chk.objects}
    t2 = back[obj.path_id].read_typetree()
    a2 = np.array(back[tex_obj.path_id].read().image)
    a2 = a2[..., 3] if a2.ndim == 3 else a2
    if len(t2["m_CharacterTable"]) != len(t["m_CharacterTable"]) or not np.array_equal(a2, alpha):
        raise SystemExit("read-back does not match what was written")
    print("wrote %s (%d bytes); asset and atlas read back identical, other objects untouched"
          % (a.out, size_b))


def preview(path, text, t, alpha, pad):
    """Draw `text` through the atlas the way the TMP shader would (no outline)."""
    import numpy as np
    from PIL import Image
    H = alpha.shape[0]
    glyphs = {g["m_Index"]: g for g in t["m_GlyphTable"]}
    chars = {c["m_Unicode"]: c["m_GlyphIndex"] for c in t["m_CharacterTable"]}
    asc = int(math.ceil(t["m_FaceInfo"]["m_AscentLine"])) + 8
    line = int(math.ceil(t["m_FaceInfo"]["m_LineHeight"]))
    width = int(sum(glyphs[chars[ord(c)]]["m_Metrics"]["m_HorizontalAdvance"]
                    for c in text if ord(c) in chars)) + 20
    canvas = np.zeros((line + 16, max(width, 20)), np.float64)
    pen = 10.0
    for c in text:
        if ord(c) not in chars:
            continue
        g = glyphs[chars[ord(c)]]
        x, y, w, h = rect(g["m_GlyphRect"])
        m = g["m_Metrics"]
        if w:
            tile = alpha[H - (y + h):H - y, x:x + w].astype(float) / 255
            cov = np.clip((tile - 0.5) * 6 + 0.5, 0, 1)
            cx = int(round(pen + math.floor(m["m_HorizontalBearingX"])))
            cy = asc - int(math.ceil(m["m_HorizontalBearingY"]))
            if cy >= 0 and cx >= 0:
                sub = canvas[cy:cy + h, cx:cx + w]
                canvas[cy:cy + h, cx:cx + w] = np.maximum(sub, cov[:sub.shape[0], :sub.shape[1]])
        pen += m["m_HorizontalAdvance"]
    Image.fromarray((255 - canvas * 255).astype(np.uint8), "L").save(path)


def main():
    utf8_stdout()
    ap = argparse.ArgumentParser(description=__doc__.split("\n")[0])
    sub = ap.add_subparsers(dest="cmd", required=True)
    p = sub.add_parser("info")
    p.add_argument("file")
    p.add_argument("--charset", action="append", choices=sorted(CHARSETS))
    p.add_argument("--text")
    p.add_argument("--sheet", help="workbook whose column C to check")
    add_common(p)
    p = sub.add_parser("verify")
    p.add_argument("file")
    p.add_argument("asset", help="TMP font asset name")
    p.add_argument("ttf", help="the asset's own source font")
    p.add_argument("--face-index", type=int, default=0)
    p.add_argument("--sample", type=int, default=60, help="how many glyphs to compare")
    add_common(p)
    p = sub.add_parser("add")
    p.add_argument("file")
    p.add_argument("asset", help="TMP font asset name")
    p.add_argument("ttf", help="font to draw the new characters from")
    p.add_argument("--face-index", type=int, default=0)
    p.add_argument("--charset", action="append", choices=sorted(CHARSETS))
    p.add_argument("--text", help="characters to add")
    p.add_argument("--sheet", help="add every character used in this workbook's column C")
    p.add_argument("--replace", action="store_true",
                   help="also re-bake requested characters the asset already has, e.g. all "
                        "Latin letters, so they match the new ones")
    p.add_argument("--preview", help="PNG showing a sample line through the new atlas")
    p.add_argument("--preview-text", help="text for --preview")
    p.add_argument("--out", required=True, help="output file (may be the input itself)")
    p.add_argument("--packer", default="auto", choices=["auto", "lz4", "lzma", "none", "original"])
    p.add_argument("--apply", action="store_true", help="write (default: dry run)")
    add_common(p)
    a = ap.parse_args()
    {"info": cmd_info, "verify": cmd_verify, "add": cmd_add}[a.cmd](a)


if __name__ == "__main__":
    main()
